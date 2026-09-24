"""Read-only behavior-event parsing and volume membership queries."""

from __future__ import annotations

import math
from dataclasses import dataclass
from numbers import Real
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook

    XLSX_AVAILABLE = True
except ImportError:  # pragma: no cover - depends on the optional extra
    load_workbook = None
    XLSX_AVAILABLE = False


@dataclass(frozen=True, slots=True)
class BehaviorEvent:
    """One behavior event on the half-open interval ``[start, stop)``."""

    label: str
    start: int
    stop: int

    def __post_init__(self) -> None:
        if not self.label:
            raise ValueError("behavior label must not be empty")
        if self.start < 0:
            raise ValueError("behavior start must be non-negative")
        if self.stop <= self.start:
            raise ValueError("behavior stop must be greater than start")


@dataclass(frozen=True, slots=True)
class BehaviorWorkbook:
    """Validated behavior data from one workbook."""

    labels: tuple[str, ...]
    events: tuple[BehaviorEvent, ...]

    @property
    def behavior_count(self) -> int:
        return len(self.labels)

    @property
    def event_count(self) -> int:
        return len(self.events)


def behavior_to_json(workbook: BehaviorWorkbook) -> dict[str, Any]:
    """Store volume intervals without depending on the source XLSX file."""
    return {
        "labels": list(workbook.labels),
        "events": [
            {"label": event.label, "start": event.start, "stop": event.stop}
            for event in workbook.events
        ],
    }


def behavior_from_json(value: Any) -> BehaviorWorkbook:
    """Validate behavior intervals in a proofread sidecar or recovery."""
    if not isinstance(value, dict) or set(value) != {"labels", "events"}:
        raise ValueError("behavior must contain labels and events")
    labels = value["labels"]
    records = value["events"]
    if (
        not isinstance(labels, list)
        or any(not isinstance(label, str) or not label for label in labels)
    ):
        raise ValueError("behavior labels must be non-empty strings")
    if not isinstance(records, list):
        raise ValueError("behavior events must be a list")
    events = []
    for index, record in enumerate(records):
        if not isinstance(record, dict) or set(record) != {"label", "start", "stop"}:
            raise ValueError(f"behavior event {index} must contain label, start, stop")
        label, start, stop = record["label"], record["start"], record["stop"]
        if (
            not isinstance(label, str)
            or label not in labels
            or type(start) is not int
            or type(stop) is not int
        ):
            raise ValueError(f"behavior event {index} has invalid label or volume")
        try:
            events.append(BehaviorEvent(label, start, stop))
        except ValueError as exc:
            raise ValueError(f"behavior event {index}: {exc}") from exc
    return BehaviorWorkbook(tuple(labels), tuple(events))


def active_behavior_labels(
    events: tuple[BehaviorEvent, ...], volume: int | None
) -> tuple[str, ...]:
    """Return active labels once each, preserving event/sheet order."""
    if volume is None or volume < 0:
        return ()
    active: list[str] = []
    seen: set[str] = set()
    for event in events:
        if (
            event.label not in seen
            and event.start <= volume < event.stop
        ):
            active.append(event.label)
            seen.add(event.label)
    return tuple(active)


def _integer_cell(value: object, *, field: str) -> int:
    if isinstance(value, bool) or not isinstance(value, Real):
        raise ValueError(f"{field} must be a number")
    numeric = float(value)
    if not math.isfinite(numeric) or not numeric.is_integer():
        raise ValueError(f"{field} must be a finite integer")
    return int(numeric)


def _event_from_values(
    label: str, values: tuple[object, ...], row_number: int
) -> BehaviorEvent:
    start_value = values[0] if values else None
    duration_value = values[1] if len(values) > 1 else None
    try:
        start = _integer_cell(start_value, field="start volume")
        duration = _integer_cell(duration_value, field="duration")
        if start < 0:
            raise ValueError("start volume must be non-negative")
        if duration <= 0:
            raise ValueError("duration must be greater than zero")
        return BehaviorEvent(label, start, start + duration)
    except ValueError as error:
        raise ValueError(
            f"Worksheet '{label}', row {row_number}: {error}"
        ) from error


def load_behavior_workbook(path: str | Path) -> BehaviorWorkbook:
    """Load behavior intervals from an XLSX workbook without modifying it."""
    source = Path(path)
    if source.suffix.casefold() != ".xlsx":
        raise ValueError("Behavior input must be an .xlsx workbook")
    if load_workbook is None:
        raise RuntimeError("Install the 'excel' extra to load XLSX workbooks")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        formula_workbook = load_workbook(
            source, read_only=True, data_only=False
        )
    except Exception:
        workbook.close()
        raise
    try:
        labels: list[str] = []
        events: list[BehaviorEvent] = []
        for sheet in workbook.worksheets:
            label = sheet.title.strip()
            if not label:
                raise ValueError(
                    f"Worksheet '{sheet.title}' has an empty behavior label"
                )
            labels.append(label)
            first_non_empty = True
            formula_sheet = formula_workbook[sheet.title]
            rows = zip(
                sheet.iter_rows(values_only=True),
                formula_sheet.iter_rows(values_only=False),
                strict=False,
            )
            for row_number, (values, formula_cells) in enumerate(
                rows, start=1
            ):
                first_two = tuple(values[:2])
                if any(
                    cell.data_type == "f" and value is None
                    for value, cell in zip(
                        first_two, formula_cells[:2], strict=False
                    )
                ):
                    raise ValueError(
                        f"Worksheet '{label}', row {row_number}: "
                        "formula has no cached numeric value"
                    )
                if all(value is None for value in first_two):
                    continue
                if first_non_empty and any(
                    isinstance(value, str) for value in first_two
                ):
                    first_non_empty = False
                    continue
                event = _event_from_values(label, first_two, row_number)
                first_non_empty = False
                events.append(event)
        return BehaviorWorkbook(tuple(labels), tuple(events))
    finally:
        workbook.close()
        formula_workbook.close()


__all__ = (
    "BehaviorEvent",
    "BehaviorWorkbook",
    "XLSX_AVAILABLE",
    "active_behavior_labels",
    "behavior_from_json",
    "behavior_to_json",
    "load_behavior_workbook",
)
