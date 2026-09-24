import pytest

from napari_worm_neuron_annotator._behavior import (
    XLSX_AVAILABLE,
    BehaviorEvent,
    active_behavior_labels,
    load_behavior_workbook,
)


def test_active_behavior_labels_use_half_open_bounds_and_sheet_order():
    events = (
        BehaviorEvent("forward", 2, 5),
        BehaviorEvent("turn", 3, 4),
        BehaviorEvent("forward", 3, 6),
    )

    assert active_behavior_labels(events, None) == ()
    assert active_behavior_labels(events, 1) == ()
    assert active_behavior_labels(events, 2) == ("forward",)
    assert active_behavior_labels(events, 3) == ("forward", "turn")
    assert active_behavior_labels(events, 4) == ("forward",)
    assert active_behavior_labels(events, 6) == ()


def test_load_behavior_workbook_rejects_non_xlsx_before_optional_import():
    with pytest.raises(ValueError, match=r"must be an \.xlsx workbook"):
        load_behavior_workbook("behavior.csv")


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl is optional")
def test_load_behavior_workbook_accepts_header_or_headerless_sheets(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "behavior.xlsx"
    workbook = Workbook()
    forward = workbook.active
    forward.title = " forward "
    forward.append(["start_volume", "duration", "ignored"])
    forward.append([2, 3, "note"])
    forward.append([None, None, "ignored"])
    forward.append([8.0, 2.0])
    turn = workbook.create_sheet("turn")
    turn.append([3, 1])
    workbook.create_sheet("pause")
    workbook.save(path)
    workbook.close()

    result = load_behavior_workbook(path)

    assert result.labels == ("forward", "turn", "pause")
    assert result.behavior_count == 3
    assert result.event_count == 3
    assert [(event.label, event.start, event.stop) for event in result.events] == [
        ("forward", 2, 5),
        ("forward", 8, 10),
        ("turn", 3, 4),
    ]


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl is optional")
@pytest.mark.parametrize(
    ("values", "message"),
    [
        ((-1, 2), "start volume must be non-negative"),
        ((1, 0), "duration must be greater than zero"),
        ((1.5, 2), "start volume must be a finite integer"),
        ((True, 2), "start volume must be a number"),
        ((1, None), "duration must be a number"),
    ],
)
def test_load_behavior_workbook_reports_invalid_data_row(
    tmp_path, values, message
):
    from openpyxl import Workbook

    path = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "turn"
    sheet.append(["start", "duration"])
    sheet.append(values)
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="Worksheet 'turn', row 2") as error:
        load_behavior_workbook(path)

    assert message in str(error.value)


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl is optional")
def test_load_behavior_workbook_does_not_hide_invalid_first_event_as_header(
    tmp_path,
):
    from openpyxl import Workbook

    path = tmp_path / "invalid-first-row.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "turn"
    sheet.append([-1, 2])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="Worksheet 'turn', row 1") as error:
        load_behavior_workbook(path)

    assert "start volume must be non-negative" in str(error.value)


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl is optional")
def test_load_behavior_workbook_rejects_formula_without_cached_value(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "turn"
    sheet.append(["start", "duration"])
    sheet.append(["=1+1", 2])
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError, match="Worksheet 'turn', row 2") as error:
        load_behavior_workbook(path)

    assert "formula has no cached numeric value" in str(error.value)
